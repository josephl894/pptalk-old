from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation


@dataclass
class Slide:
    index: int
    title: str
    body: str
    speaker_notes: str = ""


def parse(pptx_path: Path) -> list[Slide]:
    prs = Presentation(str(pptx_path))
    slides: list[Slide] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_text = ""
        body_parts: list[str] = []
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if title_shape is not None and shape == title_shape:
                title_text = text
            else:
                body_parts.append(text)

        notes_text = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes_text = notes_frame.text.strip()

        slides.append(
            Slide(
                index=i,
                title=title_text,
                body="\n".join(body_parts),
                speaker_notes=notes_text,
            )
        )
    return slides
